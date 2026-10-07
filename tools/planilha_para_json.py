"""Converte a sua planilha de orçamento (.xlsx, uma aba por ano) no arquivo JSON que o app importa em Ajustes.

Uso (no terminal, dentro da pasta do projeto):
    pip install openpyxl                      # uma vez
    python tools/planilha_para_json.py "C:/caminho/Orçamento completo.xlsx" --saida "C:/caminho/historico.json" --saldo-inicial 2026=90000

O arquivo gerado tem dados pessoais: NUNCA o envie ao GitHub (o .gitignore já bloqueia historico.json).

Regras que o conversor segue (as mesmas explicadas no docs/PEDIDOS-E-DECISOES.md):
  - células com "#" (valor já incluso na fatura do cartão) viram itens "informativos": aparecem, mas não somam;
  - o "OK" ao lado de cada mês marca o item como pago/recebido; sem "OK" fica como previsto;
  - as anotações (comentários) da célula viram os itens "valor - descrição"; se não somam o total da célula,
    entra uma linha "(sem detalhe / diferença)" para o total continuar igual ao da planilha;
  - linhas "sobrou 20XX" viram o saldo inicial do ano (não são receita);
  - abas 2017+ têm o layout "Receber / PAGAR com meses lado a lado"; 2015 e 2016 têm o layout antigo por grupos.
"""
import argparse, json, re, sys, unicodedata
import openpyxl

MES = ['janeiro', 'fevereiro', 'março', 'abril', 'maio', 'junho', 'julho', 'agosto', 'setembro', 'outubro', 'novembro', 'dezembro']


# ---------- leitura de números e anotações ----------
def num(s):
    s = s.strip()
    m = re.match(r'^(-?)\s*(?:R\$)?\s*(-?)\s*([\d][\d.,]*)$', s)
    if not m:
        return None
    neg = bool(m.group(1) or m.group(2)); d = m.group(3)
    seps = [i for i, ch in enumerate(d) if ch in '.,']
    if seps:
        last = seps[-1]
        v = float(re.sub(r'[.,]', '', d[:last]) + '.' + d[last + 1:]) if len(d) - last - 1 == 2 else float(re.sub(r'[.,]', '', d))
    else:
        v = float(d)
    return -v if neg else v

def tonum(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        return num(v.replace('#', '').strip())
    return None

def parse_note(t):
    out = []
    for line in t.splitlines():
        m = re.match(r'^(-?\s*(?:R\$)?\s*-?\s*[\d][\d.,]*)\s*[-–]\s*(.+)$', line.strip())
        if m:
            v = num(m.group(1))
            if v is not None:
                out.append([v, m.group(2).strip()])
    return out

def norm(s):
    return re.sub(r'\s+', ' ', ''.join(c for c in unicodedata.normalize('NFD', s.lower()) if unicodedata.category(c) != 'Mn')).strip()


class Conversor:
    def __init__(self, caminho):
        self.wv = openpyxl.load_workbook(caminho, data_only=True)   # valores calculados
        self.wf = openpyxl.load_workbook(caminho)                    # fórmulas e comentários
        self.cats = {}; self.data = {}; self.saldo_ini = {}

    def cat_id(self, tipo, nome, dia):
        k = (tipo, norm(nome))
        if k not in self.cats:
            self.cats[k] = {'id': f'i{len(self.cats) + 1}', 'tipo': tipo, 'nome': nome.strip(), 'dia': dia}
        elif dia and not self.cats[k]['dia']:
            self.cats[k]['dia'] = dia
        return self.cats[k]['id']

    # ---------- abas 2017 em diante ----------
    def aba_nova(self, year):
        ws, wc = self.wv[str(year)], self.wf[str(year)]
        hdr = None
        for r in range(1, 5):
            if ws.cell(r, 3).value == 'Receber':
                hdr = r
        mcols = {}
        for c in range(1, 40):
            v = ws.cell(hdr, c).value
            if isinstance(v, str) and v.strip().lower() in MES:
                mcols[MES.index(v.strip().lower())] = c
        contiguous = mcols[1] == mcols[0] + 1          # 2021: meses colados, sem coluna "OK"
        tipo = 'receber'; rows = []; tot = {'receber': None, 'pagar': None}
        for r in range(hdr + 1, ws.max_row + 1):
            n = ws.cell(r, 3).value
            if not isinstance(n, str) or not n.strip():
                continue
            nn = n.strip()
            if nn.upper() == 'PAGAR':
                tipo = 'pagar'; continue
            if nn.upper() == 'TOTAL':
                tot[tipo] = r
                if tipo == 'pagar':
                    break
                continue
            rows.append((tipo, r, nn))
        stats = {'celulas': 0, 'ajustes': 0, 'sem_ok': 0, 'cartao': 0}
        mine = {'receber': [0] * 12, 'pagar': [0] * 12}
        for tipo, r, nn in rows:
            sobrou = norm(nn).startswith('sobrou')
            dia = ws.cell(r, 2).value
            dia = int(dia) if isinstance(dia, (int, float)) and 1 <= dia <= 31 else None
            cid = None
            for m, col in mcols.items():
                raw = ws.cell(r, col).value
                v = tonum(raw); cm = wc.cell(r, col).comment
                ok = True if contiguous else str(ws.cell(r, col + 1).value or '').strip().upper() == 'OK'
                items = parse_note(cm.text) if cm else []
                info = isinstance(raw, str) and '#' in raw
                if v is None and not items:
                    continue
                v = v or 0.0
                if sobrou:
                    if v:
                        self.saldo_ini[year] = self.saldo_ini.get(year, 0) + v
                    mine[tipo][m] += v
                    continue
                diff = round(v - sum(i[0] for i in items), 2)
                if info:
                    if not items:
                        items = [[v, '(no cartão)']]
                    cid = cid or self.cat_id(tipo, nn, dia)
                    self.data[f'{cid}|{year}|{m}'] = {'items': [{'v': round(i[0], 2), 'd': i[1], 'ok': ok, 'info': True} for i in items]}
                    stats['cartao'] += 1
                    continue
                if items and abs(diff) > 0.005:
                    items.append([diff, '(sem detalhe / diferença)']); stats['ajustes'] += 1
                if not items:
                    if v == 0:
                        continue
                    items = [[v, '']]
                cid = cid or self.cat_id(tipo, nn, dia)
                if not ok and v:
                    stats['sem_ok'] += 1
                self.data[f'{cid}|{year}|{m}'] = {'items': [{'v': round(i[0], 2), 'd': i[1], 'ok': ok} for i in items]}
                stats['celulas'] += 1; mine[tipo][m] += v
        bad = []
        for t in ('receber', 'pagar'):
            if tot[t] is None:
                continue
            for m, col in mcols.items():
                sv = tonum(ws.cell(tot[t], col).value) or 0
                if abs(sv - mine[t][m]) > 1:
                    bad.append((t, MES[m][:3], round(sv, 2), round(mine[t][m], 2)))
        print(year, stats, '| diferenças contra o TOTAL da planilha:', bad[:6], '...' if len(bad) > 6 else '')

    # ---------- abas 2015 e 2016 (layout antigo, por grupos) ----------
    def aba_antiga(self, year):
        ws, wc = self.wv[str(year)], self.wf[str(year)]
        group = None; tipo = None; mine = {'receber': [0] * 12, 'pagar': [0] * 12}; stats = {'celulas': 0, 'cartao': 0}
        for r in range(1, ws.max_row + 1):
            a = ws.cell(r, 1).value; b = ws.cell(r, 2).value
            if isinstance(a, str) and a.strip():
                group = a.strip()
                if group.upper().startswith('PLANILHA'):
                    group = None
                tipo = 'receber' if group and group.upper() == 'MINHA RENDA' else 'pagar'
                continue
            if not isinstance(b, str) or not b.strip():
                continue
            nn = b.strip()
            if nn.upper() == 'TOTAIS':
                break
            if group is None or re.fullmatch(r'FIXO( \d+)?|Alguma coisa \d|Novo|Qualquer cartão', nn, re.I) and not any(tonum(ws.cell(r, c).value) for c in range(3, 15)):
                continue
            cid = None
            for m in range(12):
                col = 3 + m; raw = ws.cell(r, col).value; v = tonum(raw); cm = wc.cell(r, col).comment
                items = parse_note(cm.text) if cm else []
                if v is None and not items:
                    continue
                v = v or 0.0; info = isinstance(raw, str) and '#' in raw
                if not items:
                    if v == 0:
                        continue
                    items = [[v, '(no cartão)' if info else '']]
                elif not info:
                    diff = round(v - sum(i[0] for i in items), 2)
                    if abs(diff) > 0.005:
                        items.append([diff, '(sem detalhe / diferença)'])
                cid = cid or self.cat_id(tipo, nn, None)
                self.data[f'{cid}|{year}|{m}'] = {'items': [dict(v=round(i[0], 2), d=i[1], ok=True, **({'info': True} if info else {})) for i in items]}
                stats['celulas'] += 1
                if info:
                    stats['cartao'] += 1
                else:
                    mine[tipo][m] += v
        rr = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 2).value == 'Rendimentos'][0]
        bad = []
        for m in range(12):
            for t, row in (('receber', rr), ('pagar', rr + 1)):
                sv = tonum(ws.cell(row, 3 + m).value) or 0
                if abs(sv - mine[t][m]) > 1:
                    bad.append((t, MES[m][:3], round(sv, 2), round(mine[t][m], 2)))
        print(year, stats, '| diferenças contra Rendimentos/Gastos:', bad)

    def gerar(self, anos):
        for y in anos:
            (self.aba_antiga if y <= 2016 else self.aba_nova)(y)
        used = {k.split('|')[0] for k in self.data}
        return {'anos': list(anos), 'saldoIni': {str(k): round(v, 2) for k, v in self.saldo_ini.items()},
                'cats': [c for c in self.cats.values() if c['id'] in used], 'data': self.data}


def main():
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8')   # evita erro de acento no console do Windows
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('xlsx', help='caminho da planilha (.xlsx)')
    ap.add_argument('--saida', default='historico.json', help='arquivo JSON a gerar')
    ap.add_argument('--anos', default='2015-2026', help='intervalo de abas a ler, ex.: 2015-2026 ou 2024-2026')
    ap.add_argument('--saldo-inicial', action='append', default=[], metavar='ANO=VALOR',
                    help='saldo inicial de um ano (ex.: 2026=90000). Pode repetir. Sobrescreve o que vier das linhas "sobrou"')
    a = ap.parse_args()
    ini, fim = (int(x) for x in a.anos.split('-'))
    c = Conversor(a.xlsx)
    out = c.gerar(range(ini, fim + 1))
    for item in a.saldo_inicial:
        y, v = item.split('=')
        out['saldoIni'][str(int(y))] = float(v)
    with open(a.saida, 'w', encoding='utf-8') as f:
        json.dump(out, f, ensure_ascii=False)
    print(f"\nGerado {a.saida}: {len(out['cats'])} categorias, {len(out['data'])} células, saldos iniciais {out['saldoIni']}")
    print('Importe em Ajustes > "Importar anos anteriores (.json)". Não envie este arquivo ao GitHub.')


if __name__ == '__main__':
    sys.exit(main())
